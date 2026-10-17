"""Planilha de taxas contratadas (mantida à mão pela Bia).

Armadilhas tratadas: cabeçalho fora da linha 1, taxa ora fração (0.0259)
ora texto ('2,59%'), CNPJ ora formatado ora inteiro sem zero à esquerda,
EC vazio, duas linhas de vigência para o mesmo lojista, spread 'padrão'.
"""
import re
from dataclasses import dataclass
from datetime import date
from decimal import Decimal
from pathlib import Path

from openpyxl import load_workbook

from ..core.dinheiro import cnpj, data_br, pct_planilha

ABA_VIGENTE = "Taxas vigentes"
ABA_ANTIGA = "ANTIGO 2025 - nao usar"


@dataclass
class TaxaContratada:
    ref_origem: str                 # 'Taxas vigentes!L13'
    lojista: str
    cnpj: str
    ec: str | None
    debito: Decimal | None          # pontos percentuais
    credito_vista: Decimal | None
    credito_2_6: Decimal | None
    credito_7_12: Decimal | None
    pix: Decimal | None
    antecipacao: Decimal | None
    spread_remessa: Decimal | None  # None quando em branco ou 'padrão'
    spread_padrao: bool             # True quando em branco ou 'padrão'
    vigencia_inicio: date | None
    vigencia_fim: date | None
    vigencia_texto: str
    obs: str
    comentario: str

    def vale_em(self, dia: date) -> bool:
        if self.vigencia_inicio and dia < self.vigencia_inicio:
            return False
        if self.vigencia_fim and dia > self.vigencia_fim:
            return False
        return True


def _vigencia(texto: str) -> tuple[date | None, date | None]:
    m = re.search(r"\d{2}/\d{2}/\d{4}", texto or "")
    if not m:
        return None, None
    dia = data_br(m.group())
    if "até" in texto.lower():
        return None, dia
    return dia, None


def _texto(v) -> str:
    return "" if v is None else str(v).strip()


def ler(caminho: Path) -> list[TaxaContratada]:
    ws = load_workbook(caminho, data_only=True)[ABA_VIGENTE]
    cabecalho = next(r[0].row for r in ws.iter_rows() if _texto(r[0].value) == "Lojista")
    taxas = []
    for row in ws.iter_rows(min_row=cabecalho + 1):
        c = [cel.value for cel in row]
        if not c[0] or c[1] is None:      # linhas vazias e notas de rodapé
            continue
        spread_txt = _texto(c[9])
        spread = pct_planilha(c[9]) if spread_txt else None
        inicio, fim = _vigencia(_texto(c[10]))
        taxas.append(TaxaContratada(
            ref_origem=f"{ABA_VIGENTE}!L{row[0].row}",
            lojista=_texto(c[0]),
            cnpj=cnpj(c[1]),
            ec=_texto(c[2]) or None,
            debito=pct_planilha(c[3]),
            credito_vista=pct_planilha(c[4]),
            credito_2_6=pct_planilha(c[5]),
            credito_7_12=pct_planilha(c[6]),
            pix=pct_planilha(c[7]),
            antecipacao=pct_planilha(c[8]),
            spread_remessa=spread,
            spread_padrao=spread is None,
            vigencia_inicio=inicio,
            vigencia_fim=fim,
            vigencia_texto=_texto(c[10]),
            obs=_texto(c[11]),
            comentario=row[0].comment.text.strip() if row[0].comment else "",
        ))
    return taxas


def ler_aba_antiga(caminho: Path) -> dict[str, str]:
    """Aba 2025: só evidência (nome -> CNPJ). Nunca usar como contrato."""
    ws = load_workbook(caminho, data_only=True)[ABA_ANTIGA]
    return {_texto(r[0]): cnpj(r[1]) for r in ws.iter_rows(min_row=2, values_only=True) if r[0]}
