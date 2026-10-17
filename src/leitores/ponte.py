"""E-mail da Ponte Câmbio com o relatório mensal (.xlsx em base64 no .eml)."""
import email
import email.utils
import io
from dataclasses import dataclass
from datetime import date, datetime
from decimal import Decimal
from email import policy
from pathlib import Path

from openpyxl import load_workbook

from ..core.dinheiro import cnpj, decimal_de_celula


@dataclass
class OperacaoCambio:
    ref_origem: str              # Nº Operação (pode repetir: duplicata)
    linha: int                   # linha física na aba Operações
    data_contratacao: date
    data_liquidacao: date
    cnpj: str
    cliente: str                 # truncado em 20 caracteres pela Ponte
    moeda: str
    valor_me: Decimal
    taxa_cliente: Decimal
    ptax_ref: Decimal
    valor_brl: Decimal
    spread_brl: Decimal
    spread_pct: Decimal          # pontos percentuais (1.5007 = 1,5007%)
    spread_estimado_d1: Decimal | None
    iof_brl: Decimal
    status: str
    observacao: str


@dataclass
class EmailPonte:
    assunto: str
    data: datetime
    corpo: str
    anexo: str
    operacoes: list[OperacaoCambio]


def _data(v) -> date:
    return v.date() if isinstance(v, datetime) else v


def ler(caminho: Path) -> EmailPonte:
    with open(caminho, "rb") as f:
        msg = email.message_from_binary_file(f, policy=policy.default)
    corpo = msg.get_body(preferencelist=("plain",))
    anexo = next(p for p in msg.iter_attachments() if (p.get_filename() or "").lower().endswith(".xlsx"))
    ws = load_workbook(io.BytesIO(anexo.get_payload(decode=True)), data_only=True).worksheets[0]

    cabecalho = None
    operacoes = []
    for row in ws.iter_rows():
        c = [cel.value for cel in row]
        primeira = str(c[0] or "").strip()
        if cabecalho is None:
            if primeira.startswith("N") and "Opera" in primeira:
                cabecalho = row[0].row
            continue
        if not primeira.startswith("PC-"):   # rodapé de totais e linhas vazias
            continue
        operacoes.append(OperacaoCambio(
            ref_origem=primeira,
            linha=row[0].row,
            data_contratacao=_data(c[1]),
            data_liquidacao=_data(c[2]),
            cnpj=cnpj(c[3]),
            cliente=str(c[4]).strip(),
            moeda=str(c[5]).strip(),
            valor_me=decimal_de_celula(c[6]),
            taxa_cliente=decimal_de_celula(c[7]),
            ptax_ref=decimal_de_celula(c[8]),
            valor_brl=decimal_de_celula(c[9]),
            spread_brl=decimal_de_celula(c[10]),
            spread_pct=(decimal_de_celula(c[11]) * 100).normalize(),
            spread_estimado_d1=decimal_de_celula(c[12]),
            iof_brl=decimal_de_celula(c[13]),
            status=str(c[14] or "").strip(),
            observacao=str(c[15] or "").strip(),
        ))
    return EmailPonte(
        assunto=str(msg["subject"]),
        data=email.utils.parsedate_to_datetime(msg["date"]),
        corpo=corpo.get_content() if corpo else "",
        anexo=anexo.get_filename(),
        operacoes=operacoes,
    )
