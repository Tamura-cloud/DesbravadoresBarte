"""Leitor do e-mail da Ponte Câmbio: .eml com o relatório de operações em XLSX anexo.

Só lê e converte tipos. Não deduplica, não filtra e não calcula comissão.
"""

from __future__ import annotations

import io
import unicodedata
from email import policy
from email.message import EmailMessage
from email.parser import BytesParser
from pathlib import Path

import openpyxl
import pandas as pd

from tarvo.leitores.comum import cnpj14, data_br, dinheiro, num_br, pct
from tarvo.leitores.layout import avisos_layout

ABA = "Operações"
CABECALHO_REF = "Nº Operação"
RODAPE_TOTAL = "Total de operações"
RODAPE_SPREAD = "Spread total R$"

# Cabeçalho da planilha -> coluna do DataFrame. É o layout do arquivo da Ponte, não regra de negócio.
COLUNAS = {
    "Nº Operação": "ref_origem",
    "Data Contratação": "data_contratacao",
    "Data Liquidação": "data_liquidacao",
    "CNPJ Cliente": "cnpj",
    "Cliente": "cliente",
    "Moeda": "moeda",
    "Valor ME": "valor_me",
    "Taxa Cliente": "taxa_cliente",
    "PTAX Ref.": "ptax_ref",
    "Valor R$": "valor_brl",
    "Spread R$": "spread_brl",
    "Spread %": "spread_pct",
    "Spread Estimado R$ (D-1)": "spread_estimado_d1_brl",
    "IOF R$": "iof_brl",
    "Status Compliance": "status_compliance",
    "Observação": "observacao",
}

COLUNAS_SAIDA = ["arquivo_origem", "ref_origem", "linha_planilha", *(c for c in COLUNAS.values() if c != "ref_origem")]

DINHEIRO = ("valor_me", "valor_brl", "spread_brl", "spread_estimado_d1_brl", "iof_brl")
TAXAS = ("taxa_cliente", "ptax_ref")
DATAS = ("data_contratacao", "data_liquidacao")
TEXTOS = ("moeda", "status_compliance")


def _texto(v) -> str:
    """Texto de célula normalizado (NFC, sem espaços nas pontas) para comparar rótulos."""
    if v is None:
        return ""
    return unicodedata.normalize("NFC", str(v)).strip()


def _rotulo(v) -> str:
    """Rótulo de rodapé sem o ":" final."""
    return _texto(v).rstrip(":").strip()


def _vazia(linha: tuple) -> bool:
    return all(_texto(v) == "" for v in linha)


def _dinheiro_ou_none(v):
    n = num_br(v)
    return None if n is None else dinheiro(n)


def _anexo_xlsx(msg: EmailMessage, caminho: Path) -> tuple[str, bytes]:
    anexos = [p for p in msg.iter_attachments() if (p.get_filename() or "").lower().endswith(".xlsx")]
    if len(anexos) != 1:
        raise ValueError(f"{caminho.name}: esperado 1 anexo .xlsx, encontrados {len(anexos)}")
    return anexos[0].get_filename(), anexos[0].get_content()


def _corpo(msg: EmailMessage) -> str | None:
    parte = msg.get_body(preferencelist=("plain", "html"))
    return None if parte is None else parte.get_content()


def _linha_cabecalho(linhas: list[tuple]) -> tuple[int, dict[str, int]]:
    """Índice da linha que tem "Nº Operação" e mapa coluna_saida -> índice da célula. Falta de coluna é erro."""
    for i, linha in enumerate(linhas):
        textos = [_texto(v) for v in linha]
        if CABECALHO_REF in textos:
            indices = {}
            for j, t in enumerate(textos):
                if t in COLUNAS:
                    indices[COLUNAS[t]] = j
            faltando = [h for h, c in COLUNAS.items() if c not in indices]
            if faltando:
                raise ValueError(f"cabeçalho sem as colunas: {', '.join(faltando)}")
            return i, indices
    raise ValueError(f"cabeçalho com '{CABECALHO_REF}' não encontrado na aba '{ABA}'")


def _valor_rodape(linha: tuple):
    """Primeira célula não vazia depois do rótulo."""
    for v in linha[1:]:
        if _texto(v) != "":
            return v
    return None


def _ler_rodape(linhas: list[tuple]) -> dict:
    rodape = {"total_operacoes_rodape": None, "spread_total_rodape": None}
    for linha in linhas:
        celulas = [v for v in linha if _texto(v) != ""]
        if not celulas:
            continue
        rotulo = _rotulo(celulas[0])
        if rotulo == RODAPE_TOTAL:
            n = num_br(celulas[1]) if len(celulas) > 1 else None
            if n is not None and n != n.to_integral_value():
                raise ValueError(f"'{RODAPE_TOTAL}' não é inteiro: {celulas[1]!r}")
            rodape["total_operacoes_rodape"] = None if n is None else int(n)
        elif rotulo == RODAPE_SPREAD:
            rodape["spread_total_rodape"] = _dinheiro_ou_none(celulas[1]) if len(celulas) > 1 else None
    return rodape


def _eh_rodape(linha: tuple) -> bool:
    celulas = [v for v in linha if _texto(v) != ""]
    return bool(celulas) and _rotulo(celulas[0]) in (RODAPE_TOTAL, RODAPE_SPREAD)


def _converter(bruto: dict, onde: str) -> dict:
    reg = dict(bruto)
    reg["ref_origem"] = _texto(bruto["ref_origem"])
    if not reg["ref_origem"]:
        raise ValueError(f"{onde}: '{CABECALHO_REF}' vazio")
    for c in DATAS:
        reg[c] = data_br(bruto[c])
    for c in DINHEIRO:
        reg[c] = _dinheiro_ou_none(bruto[c])
    for c in TAXAS:
        reg[c] = num_br(bruto[c])
    reg["spread_pct"] = pct(bruto["spread_pct"])
    reg["cnpj"] = cnpj14(bruto["cnpj"])
    if reg["cnpj"] is None and _texto(bruto["cnpj"]) != "":  # vazio fica None: o ajuste da Ponte pode vir sem CNPJ
        raise ValueError(f"{onde}: CNPJ inválido: {bruto['cnpj']!r}")
    reg["cliente"] = None if bruto["cliente"] is None else str(bruto["cliente"])
    for c in TEXTOS:
        reg[c] = _texto(bruto[c]) or None
    reg["observacao"] = _texto(bruto["observacao"]) or None
    return reg


def ler_ponte(caminho_eml, avisos: list[str] | None = None) -> tuple[pd.DataFrame, dict]:
    """Lê o .eml da Ponte Câmbio. Devolve (operações, metadados do e-mail e do rodapé).

    Cada linha da aba "Operações" entre o cabeçalho e a primeira linha vazia vira uma linha do
    DataFrame, inclusive duplicatas e spreads negativos. O rodapé não é operação: vai para os metadados.
    avisos (opcional) recebe as colunas a mais ou fora de ordem (a leitura é pelo nome, R44).
    """
    caminho = Path(caminho_eml)
    with open(caminho, "rb") as f:
        msg = BytesParser(policy=policy.default).parse(f)

    nome_anexo, conteudo = _anexo_xlsx(msg, caminho)
    arquivo_origem = f"{caminho.name}#{nome_anexo}"

    # Sem read_only: nele as linhas vazias podem sumir, e precisamos delas para parar e numerar.
    wb = openpyxl.load_workbook(io.BytesIO(conteudo), data_only=True)
    try:
        if ABA not in wb.sheetnames:
            raise ValueError(f"{arquivo_origem}: aba '{ABA}' não encontrada (abas: {wb.sheetnames})")
        linhas = [tuple(r) for r in wb[ABA].iter_rows(values_only=True)]
    finally:
        wb.close()

    i_cab, indices = _linha_cabecalho(linhas)
    if avisos is not None:
        avisos += avisos_layout(f"{arquivo_origem}, aba '{ABA}'", [_texto(v) for v in linhas[i_cab]], COLUNAS)

    registros = []
    i = i_cab + 1
    while i < len(linhas) and not _vazia(linhas[i]) and not _eh_rodape(linhas[i]):
        linha = linhas[i]
        linha_planilha = i + 1  # linhas do Excel começam em 1
        bruto = {c: (linha[j] if j < len(linha) else None) for c, j in indices.items()}
        reg = _converter(bruto, f"{arquivo_origem} linha {linha_planilha}")
        reg["arquivo_origem"] = arquivo_origem
        reg["linha_planilha"] = linha_planilha
        registros.append(reg)
        i += 1

    # dtype object: sem ele o pandas troca None por NaN nas colunas de texto.
    df = pd.DataFrame(registros, columns=COLUNAS_SAIDA, dtype=object)
    df["linha_planilha"] = df["linha_planilha"].astype("int64")

    data_hora = msg["Date"].datetime if msg["Date"] is not None else None
    metadados = {
        "arquivo_origem": arquivo_origem,
        "remetente": None if msg["From"] is None else str(msg["From"]),
        "data_email": None if data_hora is None else data_hora.date(),
        "data_hora_email": data_hora,
        "assunto": None if msg["Subject"] is None else str(msg["Subject"]),
        "corpo": _corpo(msg),
        **_ler_rodape(linhas[i:]),
    }
    return df, metadados
