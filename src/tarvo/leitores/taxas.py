"""Leitor da planilha de taxas contratadas (XLSX mantido à mão pelo financeiro).

Só lê e converte tipos, sem regra de negócio: textos como "padrão", "INATIVO" ou "só remessa"
são preservados como vieram e interpretados depois.
"""

from __future__ import annotations

import re
import zipfile
from datetime import date, datetime
from pathlib import Path

import openpyxl
import pandas as pd
from openpyxl.utils.exceptions import InvalidFileException

from tarvo.leitores.comum import cnpj14, data_br, pct, sem_acento
from tarvo.leitores.layout import avisos_layout
from tarvo.leitores.rejeitadas import LinhaRejeitada

# Só esta aba. A aba histórica (2025) nunca é lida.
ABA = "Taxas vigentes"

# Cabeçalho da planilha (comparado sem acento e em maiúsculas) -> campo interno.
CABECALHOS = {
    "LOJISTA": "lojista",
    "CNPJ": "cnpj",
    "EC ORBE": "ec",
    "DEBITO": "taxa_debito",
    "CREDITO A VISTA": "taxa_credito_vista",
    "CRED. 2-6X": "taxa_parc_2_6",
    "CRED. 7-12X": "taxa_parc_7_12",
    "PIX": "taxa_pix",
    "ANTECIPACAO": "taxa_antecipacao",
    "SPREAD REMESSA": "spread",
    "VIGENCIA": "vigencia",
    "OBS": "obs",
}

CAMPOS_TAXA = ["taxa_debito", "taxa_credito_vista", "taxa_parc_2_6", "taxa_parc_7_12", "taxa_pix", "taxa_antecipacao"]

COLUNAS_SAIDA = [
    "arquivo_origem", "ref_origem", "lojista", "cnpj", "ec", *CAMPOS_TAXA,
    "spread_remessa_pct", "spread_remessa_texto", "vigencia_inicio", "vigencia_fim", "obs",
]

_DATA = r"(\d{2}/\d{2}/\d{4})"
_VIGENCIA_ATE = re.compile(rf"^ATE\s+{_DATA}$")
_VIGENCIA_A_PARTIR = re.compile(rf"^A PARTIR DE\s+{_DATA}$")
_VIGENCIA_DATA = re.compile(rf"^{_DATA}$")


def _vazio(v) -> bool:
    return v is None or (isinstance(v, str) and not v.strip())


def _vigencia(v) -> tuple[date | None, date | None]:
    """(início, fim). "até d" -> fim; "a partir de d" -> início; data simples -> início; vazio -> (None, None)."""
    if _vazio(v):
        return None, None
    if isinstance(v, (datetime, date)):
        return data_br(v), None
    texto = sem_acento(v)
    if m := _VIGENCIA_ATE.match(texto):
        return None, data_br(m.group(1))
    if m := _VIGENCIA_A_PARTIR.match(texto):
        return data_br(m.group(1)), None
    if m := _VIGENCIA_DATA.match(texto):
        return data_br(m.group(1)), None
    raise ValueError(f"vigência em formato desconhecido: {v!r}")


def _spread(v):
    """(pontos percentuais ou None, texto original ou None). Texto sem número ("padrão") fica só no texto."""
    if _vazio(v):
        return None, None
    if isinstance(v, str):
        return (pct(v) if v.strip().endswith("%") else None), v
    return pct(v), None


def _ec(v) -> str | None:
    if _vazio(v):
        return None
    if isinstance(v, float) and v.is_integer():
        v = int(v)
    return str(v).strip()


def _ec_seguro(v) -> str | None:
    try:
        return _ec(v)
    except (TypeError, ValueError):
        return None


def _cnpj(v) -> str:
    c = cnpj14(v)
    if c is None:
        raise ValueError(f"CNPJ inválido: {v!r}")
    return c


def _localizar_cabecalho(ws) -> tuple[int, dict[str, int], list[str]]:
    """Linha do cabeçalho (a que tem a célula "Lojista"), coluna de cada campo e os nomes lidos (sem acento)."""
    for row in ws.iter_rows():
        if any(sem_acento(c.value) == "LOJISTA" for c in row if isinstance(c.value, str)):
            nomes = [sem_acento(c.value) if isinstance(c.value, str) else "" for c in row]
            colunas = {}
            for c in row:
                if isinstance(c.value, str) and sem_acento(c.value) in CABECALHOS:
                    colunas[CABECALHOS[sem_acento(c.value)]] = c.column
            faltando = sorted(set(CABECALHOS.values()) - set(colunas))
            if faltando:
                raise ValueError(f"aba {ABA!r}, linha {row[0].row}: cabeçalhos ausentes: {faltando}")
            return row[0].row, colunas, nomes
    raise ValueError(f"aba {ABA!r}: célula 'Lojista' não encontrada")


def _converter(valores: dict, arquivo_origem: str, linha: int) -> dict:
    spread_pct, spread_texto = _spread(valores["spread"])
    inicio, fim = _vigencia(valores["vigencia"])
    return {
        "arquivo_origem": arquivo_origem,
        "ref_origem": f"linha {linha}",
        "lojista": valores["lojista"],
        "cnpj": _cnpj(valores["cnpj"]),
        "ec": _ec(valores["ec"]),
        **{campo: pct(valores[campo]) for campo in CAMPOS_TAXA},
        "spread_remessa_pct": spread_pct,
        "spread_remessa_texto": spread_texto,
        "vigencia_inicio": inicio,
        "vigencia_fim": fim,
        "obs": valores["obs"],
    }


def ler_taxas(caminho, avisos: list[str] | None = None,
              rejeitadas: list[LinhaRejeitada] | None = None) -> pd.DataFrame:
    """Lê a aba "Taxas vigentes". ref_origem = "linha N" do Excel. Para na primeira linha sem lojista.

    avisos (opcional) recebe as colunas a mais ou fora de ordem (a leitura é pelo nome, R44). Com rejeitadas, a linha
    com dado inválido e CNPJ legível vai para a lista e a leitura segue: o lojista daquele CNPJ fica fora do checkout
    e da remessa (R44). Sem CNPJ legível não há como saber de quem é a linha, e a planilha inteira é erro.
    """
    caminho = Path(caminho)
    try:
        wb = openpyxl.load_workbook(caminho, data_only=True)
    except (zipfile.BadZipFile, InvalidFileException, KeyError) as e:  # não é um xlsx (ou está corrompido)
        raise ValueError(f"{caminho.name}: não é um xlsx legível ({type(e).__name__}: {e})") from None
    try:
        if ABA not in wb.sheetnames:
            raise ValueError(f"{caminho.name}: aba {ABA!r} não encontrada (abas: {wb.sheetnames})")
        ws = wb[ABA]
        linha_cab, colunas, nomes = _localizar_cabecalho(ws)
        if avisos is not None:
            avisos += avisos_layout(f"{caminho.name}, aba {ABA!r}", nomes, CABECALHOS)
        registros = []
        for linha in range(linha_cab + 1, ws.max_row + 1):
            valores = {campo: ws.cell(row=linha, column=col).value for campo, col in colunas.items()}
            if _vazio(valores["lojista"]):
                break
            try:
                registros.append(_converter(valores, caminho.name, linha))
            except ValueError as e:
                erro = f"{caminho.name}, aba {ABA!r}, linha {linha}: {e}"
                cnpj = cnpj14(valores["cnpj"])
                if rejeitadas is None or cnpj is None:
                    raise ValueError(erro) from None
                rejeitadas.append(LinhaRejeitada(caminho.name, f"linha {linha}", erro, cnpj=cnpj,
                                                 ec=_ec_seguro(valores["ec"])))
    finally:
        wb.close()
    # dtype=object: mantém Decimal, date e None (sem virar float/NaN)
    return pd.DataFrame(registros, columns=COLUNAS_SAIDA, dtype=object)
