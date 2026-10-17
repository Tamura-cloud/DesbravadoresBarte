"""R27: ajuste de contestação da Ponte no relatório do mês seguinte (crédito ou débito).

O relatório de setembro não tem linha de ajuste de agosto: todas as operações têm data de contratação em
setembro e a observação é vazia, "reprocessada" ou "Taxa travada no pedido (D-1)". Não há exemplo real do
formato, então a fixture supõe o mínimo que deixa a linha rastreável: o Nº Operação e a Data Contratação da
operação original, o CNPJ e o Spread R$ do ajuste (crédito positivo, débito negativo), sem valor em moeda.

O pack de outubro é gerado em pasta temporária a partir de setembro (dados/ só é lido).
"""

import csv
import io
from datetime import datetime
from decimal import Decimal
from email import policy
from email.parser import BytesParser
from pathlib import Path

import openpyxl
import pytest

from gerar_outubro import gerar
from tarvo.fechamento import fechar, ler_entradas
from tarvo.leitores.ponte import ler_ponte

D = Decimal
RAIZ = Path(__file__).resolve().parents[1]
SETEMBRO = RAIZ / "dados" / "2026-09"
MUNDO_ABERTO = "17217888000111"
AJUSTES = (D("-61.69"), D("12.34"))  # débito e crédito de teste


def _ler(pasta: Path, nome: str) -> list[dict]:
    with open(pasta / nome, encoding="utf-8", newline="") as f:
        return list(csv.DictReader(f, delimiter=";"))


def _inserir_ajustes(eml: Path, ajustes: list[dict]) -> None:
    """Insere as linhas de ajuste depois da última operação da aba "Operações" do anexo."""
    with open(eml, "rb") as f:
        msg = BytesParser(policy=policy.default).parse(f)
    anexo = next(p for p in msg.iter_attachments() if p.get_filename().lower().endswith(".xlsx"))
    wb = openpyxl.load_workbook(io.BytesIO(anexo.get_content()))
    ws = wb["Operações"]
    for linha in ws.iter_rows():
        valores = [c.value for c in linha]
        if "Nº Operação" in valores:
            titulos = {v: c.column for v, c in zip(valores, linha) if v}
            ultima = linha[0].row
            break
    while ws.cell(ultima + 1, titulos["Nº Operação"]).value:
        ultima += 1
    for ajuste in ajustes:
        ws.insert_rows(ultima + 1)
        for titulo, valor in ajuste.items():
            c = ws.cell(ultima + 1, titulos[titulo], valor)
            c.number_format = ws.cell(ultima, titulos[titulo]).number_format
        ultima += 1
    saida = io.BytesIO()
    wb.save(saida)
    maintype, subtype = anexo.get_content_type().split("/")
    anexo.set_content(saida.getvalue(), maintype=maintype, subtype=subtype, disposition="attachment",
                      filename=anexo.get_filename())
    eml.write_bytes(msg.as_bytes())


@pytest.fixture(scope="module")
def originais():
    """Duas operações da Mundo Aberto de setembro: as que o ajuste de outubro corrige."""
    ponte = ler_entradas(SETEMBRO, "2026-09").ponte
    ops = ponte[ponte["cnpj"] == MUNDO_ABERTO].sort_values("ref_origem").head(len(AJUSTES))
    return ops.to_dict("records")


@pytest.fixture(scope="module")
def pack(originais, tmp_path_factory):
    pasta = tmp_path_factory.mktemp("dados") / "2026-10"
    casos = gerar(SETEMBRO, pasta)
    ajustes = [{
        "Nº Operação": op["ref_origem"],
        "Data Contratação": datetime.combine(op["data_contratacao"], datetime.min.time()),
        "CNPJ Cliente": op["cnpj"],
        "Cliente": op["cliente"],
        "Moeda": op["moeda"],
        "Spread R$": float(valor),
        "Observação": f"Ajuste de contestação - {'crédito' if valor > 0 else 'débito'}",
    } for op, valor in zip(originais, AJUSTES)]
    _inserir_ajustes(casos["ponte"]["arquivo"], ajustes)
    return {"pasta": pasta, "eml": casos["ponte"]["arquivo"]}


@pytest.fixture(scope="module")
def saida(pack, tmp_path_factory):
    pasta = tmp_path_factory.mktemp("saida")
    fechar(pack["pasta"], pasta, "2026-10", usar_ia=False)  # ErroFechamento aqui já reprova
    return pasta


def test_leitor_le_o_ajuste_com_a_referencia_da_operacao_original(pack, originais):
    df, _ = ler_ponte(pack["eml"])
    ajustes = df[df["ref_origem"].isin([op["ref_origem"] for op in originais])].to_dict("records")
    assert [a["ref_origem"] for a in ajustes] == [op["ref_origem"] for op in originais]
    for a, op, valor in zip(ajustes, originais, AJUSTES):
        assert a["data_contratacao"] == op["data_contratacao"] and a["cnpj"] == MUNDO_ABERTO
        assert a["spread_brl"] == valor and a["valor_me"] is None and a["spread_pct"] is None
        assert a["arquivo_origem"].startswith(pack["eml"].name) and a["linha_planilha"] > 0


def test_ajuste_nao_vira_operacao_nem_comissao_e_fica_rastreavel(saida, originais):
    rastro = _ler(saida, "rastro.csv")
    for op, valor in zip(originais, AJUSTES):
        ref = op["ref_origem"]
        linhas = [l for l in rastro if l["ref_origem"] == ref]
        assert [(l["id_saida"], D(l["valor"])) for l in linhas] == [(f"PEN-FORA-REM-{ref}", valor)]
    pendencias = {l["id_saida"]: l for l in _ler(saida, "pendencias.csv")}
    for op, valor in zip(originais, AJUSTES):
        pen = pendencias[f"PEN-FORA-REM-{op['ref_origem']}"]
        assert pen["ref_origem"] == op["ref_origem"] and D(pen["valor_brl"]) == valor
        assert pen["quem_decide"] == "Financeiro (Helena)"
    refs = {op["ref_origem"] for op in originais}
    for nome in ("comissoes_por_lojista.csv", "notas_a_emitir.csv", "divergencias_taxa.csv"):
        ids = {l["id_saida"] for l in _ler(saida, nome)}
        assert not [l for l in rastro if l["id_saida"] in ids and l["ref_origem"] in refs], nome
