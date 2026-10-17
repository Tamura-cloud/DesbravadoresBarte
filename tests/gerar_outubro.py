"""Gera um data pack de "outubro" a partir do de setembro, no mesmo formato, para ensaiar o fechamento.

Uso: python tests/gerar_outubro.py <pasta de setembro> <pasta de destino>

Os originais só são lidos. No destino:
- nomes de arquivo com "set2026" viram "out2026"; a planilha de taxas e o plano de contas são copiados
  byte a byte (contrato e plano não mudam com o mês);
- toda data anda +1 mês (setembro -> outubro, agosto -> setembro), cada uma no formato do seu arquivo:
  dd/mm/aaaa e mm/aaaa em texto (Orbe, Aurora, corpo e assunto do e-mail, células de texto do anexo),
  datas do Excel no anexo da Ponte, AAAAMMDD nas tags DT* do OFX, cabeçalho Date do e-mail.
  O dia se mantém; o último dia do mês vira o último dia do mês seguinte (30/09 -> 31/10);
- identificadores com o mês embutido andam junto: Nº Operação "PC-MMAA-nnnn", FITID "MRDAAAAMMnnnn"
  e o nome do anexo "..._MMAAAA.xlsx";
- quatro casos novos, um por arquivo (constantes CASO_* abaixo):
  (a) Orbe: uma venda de um EC que não está na planilha;
  (b) Aurora: "PIX ENVIADO GRAFICA RAPIDA LTDA" de -1.250,00, com o saldo corrido recalculado;
  (c) Meridiano: um "PIX RECEB" de um pagador fora da lista, com o LEDGERBAL atualizado;
  (d) Ponte: uma operação de um CNPJ fora da planilha, com rodapé e aba Resumo atualizados.

Os valores dos casos novos são de teste e ficam só aqui (tests/), nunca em src/ ou config/.
"""

from __future__ import annotations

import calendar
import io
import re
import shutil
import sys
from datetime import date, datetime
from decimal import ROUND_DOWN, Decimal
from email import policy
from email.parser import BytesParser
from email.utils import format_datetime
from pathlib import Path

import openpyxl

CENTAVO = Decimal("0.01")
RAIZ = Path(__file__).resolve().parents[1]
PROTEGIDAS = (RAIZ / "dados", RAIZ / "saida", RAIZ / "contestacoes")  # nunca recebem o pack simulado

# Tipos que podem vir em vários arquivos no mesmo formato (config/entradas.yaml, multiplos: true). O caso novo
# vai no primeiro em ordem de nome (o relatório); os outros (ajustes) só têm as datas deslocadas.
MULTIPLOS = ("orbe",)

# (a) Orbe: venda de crédito à vista de um EC fora da planilha e fora de ec_sem_codigo.
CASO_ORBE = {
    "id_transacao": "TX88999001",
    "dia": 15,
    "cod_estabelecimento": "1099300457",
    "nome_estabelecimento": "ESCOLA NOVA ERA IDIOMAS",
    "tipo_lancamento": "VENDA",
    "modalidade": "CREDITO A VISTA",
    "qtd_parcelas": "1",
    "bandeira": "VISA",
    "valor_bruto": Decimal("2450.00"),
    "taxa_aplicada_pct": Decimal("2.69"),
}

# (b) Aurora: PIX para fornecedor que não está em config/favorecidos.yaml.
CASO_AURORA = {"dia": 15, "descricao": "PIX ENVIADO GRAFICA RAPIDA LTDA", "valor": Decimal("-1250.00")}

# (c) Meridiano: PIX de pagador que não está em config/pagadores_meridiano.yaml.
CASO_MERIDIANO = {"dia": 20, "memo": "PIX RECEB INSTITUTO SAKURA IDIOMAS", "valor": Decimal("25000.00"),
                  "checknum": "512734"}

# (d) Ponte: operação de CNPJ que não está na planilha de taxas. Spread de 1,5% sobre a PTAX.
CASO_PONTE = {
    "sufixo": "9001",
    "dia_contratacao": 14,
    "dia_liquidacao": 15,
    "cnpj_base": "41928553",   # os 2 dígitos verificadores são calculados
    "cliente": "ESCOLA NOVA FRONTEIR",
    "moeda": "USD",
    "valor_me": Decimal("3000.00"),
    "ptax_ref": Decimal("5.4000"),
    "taxa_cliente": Decimal("5.4810"),
    "spread_estimado_d1_brl": Decimal("236.40"),
    "iof_pct": Decimal("0.38"),
    "status": "APROVADO",
}

_RE_DATA_TEXTO = re.compile(
    r"(?<![\d/])(\d{2})/(\d{2})/(\d{4})(?![\d/])"   # dd/mm/aaaa
    r"|(?<![\d/])(\d{2})/(\d{4})(?![\d/])"          # mm/aaaa
)
_RE_DT_OFX = re.compile(r"<(DTSERVER|DTSTART|DTEND|DTPOSTED|DTASOF)>(\d{4})(\d{2})(\d{2})")
_RE_FITID = re.compile(r"<FITID>MRD(\d{4})(\d{2})(\d{4})")
_RE_OPERACAO = re.compile(r"^PC-(\d{2})(\d{2})-(\d{4})$")
_RE_ANEXO = re.compile(r"_(\d{2})(\d{4})\.xlsx$", re.I)


# ---------------------------------------------------------------- datas

def mes_seguinte(ano: int, mes: int) -> tuple[int, int]:
    return (ano + 1, 1) if mes == 12 else (ano, mes + 1)


def deslocar(d: date) -> date:
    """+1 mês, mantendo o dia; o último dia do mês vira o último dia do mês seguinte."""
    ano, mes = mes_seguinte(d.year, d.month)
    ultimo = calendar.monthrange(ano, mes)[1]
    dia = ultimo if d.day == calendar.monthrange(d.year, d.month)[1] else min(d.day, ultimo)
    if isinstance(d, datetime):
        return d.replace(year=ano, month=mes, day=dia)
    return date(ano, mes, dia)


def _mes_valido(mes: str, ano: str) -> bool:
    return 1 <= int(mes) <= 12 and 2000 <= int(ano) <= 2099


def deslocar_texto(texto: str) -> str:
    """Desloca toda data dd/mm/aaaa e toda referência mm/aaaa do texto."""
    def trocar(m: re.Match) -> str:
        if m.group(1):
            dia, mes, ano = m.group(1), m.group(2), m.group(3)
            if not _mes_valido(mes, ano) or not 1 <= int(dia) <= calendar.monthrange(int(ano), int(mes))[1]:
                return m.group(0)
            return deslocar(date(int(ano), int(mes), int(dia))).strftime("%d/%m/%Y")
        mes, ano = m.group(4), m.group(5)
        if not _mes_valido(mes, ano):
            return m.group(0)
        a, n = mes_seguinte(int(ano), int(mes))
        return f"{n:02d}/{a}"
    return _RE_DATA_TEXTO.sub(trocar, texto)


def _data_br(s: str) -> date:
    return datetime.strptime(s, "%d/%m/%Y").date()


# ---------------------------------------------------------------- números

def _br(v: Decimal) -> str:
    """Decimal para o formato do extrato: "-1.250,00"."""
    texto = f"{abs(v):,.2f}".replace(",", "X").replace(".", ",").replace("X", ".")
    return ("-" if v < 0 else "") + texto


def _num_br(s: str) -> Decimal:
    return Decimal(s.replace(".", "").replace(",", "."))


def _virgula(v: Decimal) -> str:
    """Decimal para o formato da Orbe: "2450,00" (sem milhar)."""
    return format(v, "f").replace(".", ",")


def cnpj_com_dv(base8: str) -> str:
    """CNPJ de 14 dígitos (matriz 0001) com os dígitos verificadores."""
    numero = base8 + "0001"
    for pesos in ([5, 4, 3, 2, 9, 8, 7, 6, 5, 4, 3, 2], [6, 5, 4, 3, 2, 9, 8, 7, 6, 5, 4, 3, 2]):
        resto = sum(int(d) * p for d, p in zip(numero, pesos)) % 11
        numero += "0" if resto < 2 else str(11 - resto)
    return numero


def _cnpj_formatado(c: str) -> str:
    return f"{c[:2]}.{c[2:5]}.{c[5:8]}/{c[8:12]}-{c[12:]}"


# ---------------------------------------------------------------- arquivos

def nome_outubro(nome: str) -> str:
    return nome.replace("set2026", "out2026")


def _linhas_com_fim(texto: str) -> list[str]:
    return texto.splitlines(keepends=True)


def _fim_de_linha(texto: str) -> str:
    return "\r\n" if "\r\n" in texto else "\n"


def gerar_orbe(origem: Path, destino: Path, com_caso: bool = True) -> dict:
    """CSV da Orbe (ASCII, CRLF): desloca data_movimento e data_venda_original e, com com_caso, insere o caso (a).

    Sem o caso (arquivos de ajuste da Orbe depois do relatório, O05), só desloca as datas e devolve {}.
    """
    texto = origem.read_bytes().decode("utf-8")
    if '"' in texto:
        raise ValueError(f"{origem.name}: campo entre aspas; o gerador separa por ';' sem csv")
    fim = _fim_de_linha(texto)
    linhas = [l.rstrip("\r\n") for l in _linhas_com_fim(texto)]
    cabecalho = linhas[0].split(";")
    i_mov, i_orig = cabecalho.index("data_movimento"), cabecalho.index("data_venda_original")
    i_id = cabecalho.index("id_transacao")

    saida, ids = [linhas[0]], set()
    for linha in linhas[1:]:
        campos = linha.split(";")
        for i in (i_mov, i_orig):
            if campos[i]:
                campos[i] = deslocar(_data_br(campos[i])).strftime("%d/%m/%Y")
        ids.add(campos[i_id])
        saida.append(";".join(campos))

    if not com_caso:
        destino.write_bytes((fim.join(saida) + fim).encode("utf-8"))
        return {}
    caso = CASO_ORBE
    if caso["id_transacao"] in ids:
        raise ValueError(f"id_transacao {caso['id_transacao']} já existe na Orbe")
    datas = [_data_br(l.split(";")[i_mov]) for l in saida[1:]]
    dia = date(datas[0].year, datas[0].month, caso["dia"])
    taxa = (caso["valor_bruto"] * caso["taxa_aplicada_pct"] / 100).quantize(CENTAVO, rounding=ROUND_DOWN)  # a Orbe trunca
    valores = {
        "id_transacao": caso["id_transacao"], "data_movimento": dia.strftime("%d/%m/%Y"),
        "cod_estabelecimento": caso["cod_estabelecimento"], "nome_estabelecimento": caso["nome_estabelecimento"],
        "tipo_lancamento": caso["tipo_lancamento"], "modalidade": caso["modalidade"],
        "qtd_parcelas": caso["qtd_parcelas"], "bandeira": caso["bandeira"],
        "valor_bruto": _virgula(caso["valor_bruto"]), "taxa_aplicada_pct": _virgula(caso["taxa_aplicada_pct"]),
        "valor_taxa": _virgula(taxa), "valor_liquido": _virgula(caso["valor_bruto"] - taxa),
        "id_referencia": "", "data_venda_original": "",
    }
    nova = ";".join(valores[c] for c in cabecalho)
    pos = 1 + max(i for i, d in enumerate(datas) if d <= dia) + 1  # depois da última linha do dia
    saida.insert(pos, nova)

    destino.write_bytes((fim.join(saida) + fim).encode("utf-8"))
    return {"ref_origem": caso["id_transacao"], "cod_estabelecimento": caso["cod_estabelecimento"],
            "valor_taxa": taxa, "valor_bruto": caso["valor_bruto"]}


def gerar_aurora(origem: Path, destino: Path) -> dict:
    """CSV do Aurora (UTF-8): desloca as datas do texto, insere o caso (b) e recalcula o saldo corrido."""
    texto = origem.read_bytes().decode("utf-8")
    fim = _fim_de_linha(texto)
    linhas = [deslocar_texto(l.rstrip("\r\n")) for l in _linhas_com_fim(texto)]

    def lancamento(campos: list[str]) -> bool:
        return bool(re.fullmatch(r"\d{2}/\d{2}/\d{4}", campos[0])) and campos[1].strip().upper() != "SALDO ANTERIOR"

    campos_linhas = [l.split(";") for l in linhas]
    i_lanc = [i for i, c in enumerate(campos_linhas) if lancamento(c)]
    primeiro = _data_br(campos_linhas[i_lanc[0]][0])
    dia = date(primeiro.year, primeiro.month, CASO_AURORA["dia"])
    pos = max(i for i in i_lanc if _data_br(campos_linhas[i][0]) <= dia) + 1
    nova = [dia.strftime("%d/%m/%Y"), CASO_AURORA["descricao"], "", _br(CASO_AURORA["valor"]), ""]
    campos_linhas.insert(pos, nova)

    saldo = None
    for n, campos in enumerate(campos_linhas):
        rotulo = campos[1].strip().upper() if len(campos) > 1 else ""
        if rotulo == "SALDO ANTERIOR":
            saldo = _num_br(campos[4])
        elif rotulo == "SALDO FINAL":
            campos[4] = _br(saldo)
        elif lancamento(campos):
            saldo += _num_br(campos[3])
            if n < pos and _num_br(campos[4]) != saldo:
                raise ValueError(f"{origem.name}: saldo do original não fecha na linha {n + 1}")
            campos[4] = _br(saldo)

    destino.write_bytes((fim.join(";".join(c) for c in campos_linhas) + fim).encode("utf-8"))
    return {"ref_origem": f"L{pos + 1}", "valor": CASO_AURORA["valor"], "data": dia}


def _fitid_seguinte(m: re.Match) -> str:
    a, n = mes_seguinte(int(m.group(1)), int(m.group(2)))
    return f"<FITID>MRD{a}{n:02d}{m.group(3)}"


def _dt_seguinte(m: re.Match) -> str:
    d = deslocar(date(int(m.group(2)), int(m.group(3)), int(m.group(4))))
    return f"<{m.group(1)}>{d.strftime('%Y%m%d')}"


def gerar_meridiano(origem: Path, destino: Path) -> dict:
    """OFX do Meridiano (cp1252, CRLF): desloca as tags DT* e o FITID, insere o caso (c) e ajusta o LEDGERBAL."""
    texto = origem.read_bytes().decode("cp1252")
    fim = _fim_de_linha(texto)
    texto = _RE_FITID.sub(_fitid_seguinte, _RE_DT_OFX.sub(_dt_seguinte, texto))

    blocos = list(re.finditer(r"<STMTTRN>.*?</STMTTRN>" + re.escape(fim), texto, re.S))
    datas = [re.search(r"<DTPOSTED>(\d{8})", b.group(0)).group(1) for b in blocos]
    fitids = [re.search(r"<FITID>MRD(\d{6})(\d{4})", b.group(0)).groups() for b in blocos]
    mes = fitids[0][0]
    dia = date(int(mes[:4]), int(mes[4:]), CASO_MERIDIANO["dia"])
    i = next((k for k, d in enumerate(datas) if d > dia.strftime("%Y%m%d")), len(blocos))
    anterior = int(fitids[i - 1][1]) if i > 0 else 0
    usados = {int(s) for _, s in fitids}
    numero = next(n for n in range(anterior + 1, 10000) if n not in usados)
    fitid = f"MRD{mes}{numero:04d}"
    bloco = fim.join([
        "<STMTTRN>", "<TRNTYPE>CREDIT", f"<DTPOSTED>{dia.strftime('%Y%m%d')}120000[-3:BRT]",
        f"<TRNAMT>{format(CASO_MERIDIANO['valor'], 'f')}", f"<FITID>{fitid}",
        f"<CHECKNUM>{CASO_MERIDIANO['checknum']}", f"<MEMO>{CASO_MERIDIANO['memo']}", "</STMTTRN>",
    ]) + fim
    corte = blocos[i].start() if i < len(blocos) else texto.index("</BANKTRANLIST>")
    texto = texto[:corte] + bloco + texto[corte:]

    def saldo(m: re.Match) -> str:
        return f"{m.group(1)}{format(Decimal(m.group(2)) + CASO_MERIDIANO['valor'], 'f')}"
    texto, n = re.subn(r"(<LEDGERBAL>\s*<BALAMT>)(-?\d+\.\d{2})", saldo, texto)
    if n != 1:
        raise ValueError(f"{origem.name}: LEDGERBAL/BALAMT não encontrado")

    destino.write_bytes(texto.encode("cp1252"))
    return {"ref_origem": fitid, "valor": CASO_MERIDIANO["valor"], "data": dia}


def _texto_celula(v):
    return deslocar_texto(v) if isinstance(v, str) else v


def _operacao_seguinte(v):
    m = _RE_OPERACAO.match(v) if isinstance(v, str) else None
    if not m:
        return v
    mes, aa = int(m.group(1)), int(m.group(2))
    a, n = mes_seguinte(2000 + aa, mes)
    return f"PC-{n:02d}{a % 100:02d}-{m.group(3)}"


def gerar_anexo_ponte(conteudo: bytes) -> tuple[bytes, dict]:
    """XLSX da Ponte: desloca datas e Nº Operação, insere o caso (d) e atualiza rodapé e Resumo."""
    wb = openpyxl.load_workbook(io.BytesIO(conteudo))
    for ws in wb.worksheets:
        for linha in ws.iter_rows():
            for c in linha:
                if isinstance(c.value, (datetime, date)):
                    c.value = deslocar(c.value)
                elif isinstance(c.value, str):
                    c.value = _operacao_seguinte(_texto_celula(c.value))

    ws = wb["Operações"]
    titulos = None
    for linha in ws.iter_rows():
        valores = [c.value for c in linha]
        if "Nº Operação" in valores:
            titulos = {v: c.column for v, c in zip(valores, linha) if v}
            linha_cab = linha[0].row
            break
    col = titulos.__getitem__
    ultima = linha_cab
    while ws.cell(ultima + 1, col("Nº Operação")).value:
        ultima += 1
    refs = {ws.cell(r, col("Nº Operação")).value for r in range(linha_cab + 1, ultima + 1)}
    exemplo = ws.cell(ultima, col("Nº Operação")).value
    ref = f"{exemplo[:8]}{CASO_PONTE['sufixo']}"
    if ref in refs:
        raise ValueError(f"Nº Operação {ref} já existe na Ponte")
    contratacao = ws.cell(ultima, col("Data Contratação")).value

    caso = CASO_PONTE
    cnpj = cnpj_com_dv(caso["cnpj_base"])
    valor_brl = (caso["valor_me"] * caso["taxa_cliente"]).quantize(CENTAVO)
    spread_brl = (caso["valor_me"] * (caso["taxa_cliente"] - caso["ptax_ref"])).quantize(CENTAVO)
    spread_pct = spread_brl / (caso["valor_me"] * caso["ptax_ref"])
    iof = (valor_brl * caso["iof_pct"] / 100).quantize(CENTAVO)
    nova = {
        "Nº Operação": ref,
        "Data Contratação": contratacao.replace(day=caso["dia_contratacao"]),
        "Data Liquidação": contratacao.replace(day=caso["dia_liquidacao"]),
        "CNPJ Cliente": _cnpj_formatado(cnpj),
        "Cliente": caso["cliente"],
        "Moeda": caso["moeda"],
        "Valor ME": float(caso["valor_me"]),
        "Taxa Cliente": float(caso["taxa_cliente"]),
        "PTAX Ref.": float(caso["ptax_ref"]),
        "Valor R$": float(valor_brl),
        "Spread R$": float(spread_brl),
        "Spread %": float(round(spread_pct, 6)),
        "Spread Estimado R$ (D-1)": float(caso["spread_estimado_d1_brl"]),
        "IOF R$": float(iof),
        "Status Compliance": caso["status"],
        "Observação": None,
    }
    ws.insert_rows(ultima + 1)
    for titulo, coluna in titulos.items():
        c = ws.cell(ultima + 1, coluna, nova[titulo])
        c.number_format = ws.cell(ultima, coluna).number_format

    def somar(ws_, rotulo: str, coluna_valor: int, delta) -> None:
        for linha in ws_.iter_rows():
            if isinstance(linha[0].value, str) and linha[0].value.strip().rstrip(":") == rotulo:
                c = linha[coluna_valor - 1]
                c.value = round(c.value + delta, 2) if isinstance(c.value, float) else c.value + delta
                return
        raise ValueError(f"anexo da Ponte: linha {rotulo!r} não encontrada")

    somar(ws, "Total de operações", 2, 1)
    somar(ws, "Spread total R$", 2, float(spread_brl))
    resumo = wb["Resumo"]
    cab = {c.value: c.column for c in resumo[1]}
    somar(resumo, caso["moeda"], cab["Qtd"], 1)
    somar(resumo, caso["moeda"], cab["Volume ME"], float(caso["valor_me"]))
    somar(resumo, caso["moeda"], cab["Spread R$"], float(spread_brl))

    saida = io.BytesIO()
    wb.save(saida)
    return saida.getvalue(), {
        "ref_origem": ref, "cnpj": cnpj, "spread_brl": spread_brl, "valor_brl": valor_brl, "iof_brl": iof,
        "data_contratacao": nova["Data Contratação"].date(),
    }


def gerar_ponte(origem: Path, destino: Path) -> dict:
    """E-mail da Ponte: assunto, Date e corpo deslocados; anexo .xlsx reconstruído dentro do .eml."""
    with open(origem, "rb") as f:
        msg = BytesParser(policy=policy.default).parse(f)

    msg.replace_header("Subject", deslocar_texto(str(msg["Subject"])))
    msg.replace_header("Date", format_datetime(deslocar(msg["Date"].datetime)))

    corpo = msg.get_body(preferencelist=("plain",))
    corpo.set_content(deslocar_texto(corpo.get_content()), charset="utf-8", cte="quoted-printable")

    anexos = [p for p in msg.iter_attachments() if (p.get_filename() or "").lower().endswith(".xlsx")]
    if len(anexos) != 1:
        raise ValueError(f"{origem.name}: esperado 1 anexo .xlsx, achei {len(anexos)}")
    anexo = anexos[0]
    nome = anexo.get_filename()

    def nome_seguinte(m: re.Match) -> str:
        a, n = mes_seguinte(int(m.group(2)), int(m.group(1)))
        return f"_{n:02d}{a}.xlsx"
    novo_nome = _RE_ANEXO.sub(nome_seguinte, nome)
    conteudo, info = gerar_anexo_ponte(anexo.get_content())
    maintype, subtype = anexo.get_content_type().split("/")
    anexo.set_content(conteudo, maintype=maintype, subtype=subtype, disposition="attachment", filename=novo_nome)

    destino.write_bytes(msg.as_bytes())
    return {**info, "anexo": novo_nome, "assunto": str(msg["Subject"])}


# ---------------------------------------------------------------- pack

GERADORES = {
    "*orbe*.csv": ("orbe", gerar_orbe),
    "*aurora*.csv": ("aurora", gerar_aurora),
    "*.ofx": ("meridiano", gerar_meridiano),
    "*.eml": ("ponte", gerar_ponte),
}
COPIAS = ("*taxas*.xlsx", "*plano_de_contas*.csv")


def gerar(origem, destino) -> dict:
    """Gera o pack de outubro em destino. Devolve, por arquivo, o caso novo inserido e o arquivo gerado."""
    origem, destino = Path(origem), Path(destino)
    if origem.resolve() == destino.resolve():
        raise ValueError("destino não pode ser a pasta de origem")
    for protegida in PROTEGIDAS:
        if destino.resolve().is_relative_to(protegida):
            raise ValueError(f"destino {destino} está dentro de {protegida}: o outubro simulado vai para uma "
                             "pasta temporária, fora do repositório (dados/ é só leitura e saida/ é só setembro)")
    destino.mkdir(parents=True, exist_ok=True)
    casos: dict = {}
    for padrao, (tipo, gerador) in GERADORES.items():
        arquivos = sorted(origem.glob(padrao), key=lambda p: p.name)
        if len(arquivos) != 1 and not (tipo in MULTIPLOS and arquivos):
            raise ValueError(f"{padrao}: esperado 1 arquivo em {origem}, achei {len(arquivos)}")
        alvo = destino / nome_outubro(arquivos[0].name)
        casos[tipo] = {**gerador(arquivos[0], alvo), "arquivo": alvo}
        for outro in arquivos[1:]:  # ajustes da Orbe (O05): mesmas datas deslocadas, sem caso novo
            gerador(outro, destino / nome_outubro(outro.name), com_caso=False)
    for padrao in COPIAS:
        for arquivo in origem.glob(padrao):
            shutil.copyfile(arquivo, destino / nome_outubro(arquivo.name))
    return casos


if __name__ == "__main__":
    if len(sys.argv) != 3:
        sys.exit("uso: python tests/gerar_outubro.py <pasta de setembro> <pasta de destino>")
    for tipo, caso in gerar(sys.argv[1], sys.argv[2]).items():
        print(f"{tipo}: {caso['arquivo'].name} (caso novo: {caso['ref_origem']})")
