from datetime import date, datetime
from decimal import Decimal
from pathlib import Path

import openpyxl
import pytest

from tarvo.cli import localizar_pela_config
from tarvo.leitores.taxas import COLUNAS_SAIDA, ler_taxas

ENTRADA = Path(__file__).resolve().parents[1] / "dados" / "2026-09"

CABECALHO = ["Lojista", "CNPJ", "EC Orbe", "Débito", "Crédito à vista", "Créd. 2-6x", "Créd. 7-12x",
             "Pix", "Antecipação", "Spread remessa", "Vigência", "Obs"]


@pytest.fixture(scope="module")
def taxas():
    caminho = localizar_pela_config(ENTRADA)["taxas"]
    return ler_taxas(caminho)


def _uma(taxas, inicio_nome):
    linhas = taxas[taxas["lojista"].str.startswith(inicio_nome)]
    assert len(linhas) == 1, inicio_nome
    return linhas.iloc[0]


def test_colunas(taxas):
    assert list(taxas.columns) == COLUNAS_SAIDA


def test_25_linhas(taxas):
    assert len(taxas) == 25
    assert taxas["ref_origem"].is_unique
    assert not taxas["lojista"].str.contains("spread em branco|atualiza|TABELA", case=False).any()


def test_cnpj_14_digitos(taxas):
    assert all(isinstance(c, str) and len(c) == 14 and c.isdigit() for c in taxas["cnpj"])
    assert _uma(taxas, "Fluente")["cnpj"] == "03229837000103"
    assert _uma(taxas, "Campus Global")["cnpj"] == "08780784000120"


def test_taxa_em_texto(taxas):
    oxford = _uma(taxas, "Oxford Prep")
    assert oxford["taxa_credito_vista"] == Decimal("2.59")
    assert isinstance(oxford["taxa_credito_vista"], Decimal)


def test_rota_global(taxas):
    rota = _uma(taxas, "Rota Global")
    assert rota["taxa_credito_vista"] == Decimal("2.49")
    assert rota["spread_remessa_pct"] == Decimal("1.5")
    assert rota["spread_remessa_texto"] is None


def test_canguru_spread_padrao_so_preservado(taxas):
    canguru = _uma(taxas, "Canguru")
    assert canguru["spread_remessa_pct"] is None
    assert canguru["spread_remessa_texto"] == "padrão"


def test_bridge_duas_vigencias(taxas):
    bridge = taxas[taxas["cnpj"] == "22224579000109"]
    assert len(bridge) == 2
    assert set(bridge["ec"]) == {"1075624862"}
    ate = bridge[bridge["vigencia_fim"].notna()].iloc[0]
    desde = bridge[bridge["vigencia_inicio"].notna()].iloc[0]
    assert ate["vigencia_fim"] == date(2026, 9, 14) and ate["vigencia_inicio"] is None
    assert desde["vigencia_inicio"] == date(2026, 9, 15) and desde["vigencia_fim"] is None


def test_tipos(taxas):
    for campo in ["taxa_debito", "taxa_credito_vista", "taxa_parc_2_6", "taxa_parc_7_12", "taxa_pix",
                  "taxa_antecipacao", "spread_remessa_pct"]:
        assert all(v is None or isinstance(v, Decimal) for v in taxas[campo]), campo
    assert all(v is None or isinstance(v, str) for v in taxas["ec"])
    assert all(v is None or type(v) is date for v in taxas["vigencia_inicio"])
    assert all(v is None or type(v) is date for v in taxas["vigencia_fim"])
    assert all(r.startswith("linha ") for r in taxas["ref_origem"])


def test_textos_preservados_sem_interpretar(taxas):
    assert _uma(taxas, "Globo Trotter")["obs"] == "INATIVO desde jul/26"
    ponte = _uma(taxas, "Ponte A")
    assert ponte["obs"] == "só remessa"
    assert ponte["ec"] is None and ponte["taxa_debito"] is None and ponte["spread_remessa_pct"] is None


def _planilha(tmp_path, linhas_dados, linha_cabecalho=3):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Taxas vigentes"
    ws.cell(row=1, column=1, value="TÍTULO")
    for col, nome in enumerate(CABECALHO, start=1):
        ws.cell(row=linha_cabecalho, column=col, value=nome)
    for i, valores in enumerate(linhas_dados, start=linha_cabecalho + 1):
        for col, v in enumerate(valores, start=1):
            ws.cell(row=i, column=col, value=v)
    antiga = wb.create_sheet("ANTIGO 2025 - nao usar")
    for col, nome in enumerate(CABECALHO, start=1):
        antiga.cell(row=1, column=col, value=nome)
    antiga.append(["Não usar Ltda", "11111111000111", "999", 0.05])
    caminho = tmp_path / "taxas_teste.xlsx"
    wb.save(caminho)
    return caminho


def test_cabecalho_fora_da_linha_5_e_para_na_primeira_vazia(tmp_path):
    caminho = _planilha(tmp_path, [
        ["Loja A", "1.234.567/0001-00", 1234567890, 0.0149, "2,59%", None, None, None, None, "1,5%",
         datetime(2026, 1, 1), None],
        ["Loja B", 1234567000100, None, None, None, None, None, None, None, "padrão", "até 30/09/2026", "obs"],
        [None],
        ["Depois do vazio", "22222222000122"],
    ], linha_cabecalho=3)
    df = ler_taxas(caminho)
    assert list(df["lojista"]) == ["Loja A", "Loja B"]
    assert list(df["ref_origem"]) == ["linha 4", "linha 5"]
    a, b = df.iloc[0], df.iloc[1]
    assert a["cnpj"] == "01234567000100" and a["ec"] == "1234567890"
    assert a["taxa_debito"] == Decimal("1.49") and a["taxa_credito_vista"] == Decimal("2.59")
    assert a["spread_remessa_pct"] == Decimal("1.5") and a["spread_remessa_texto"] == "1,5%"
    assert a["vigencia_inicio"] == date(2026, 1, 1) and a["vigencia_fim"] is None
    assert b["cnpj"] == "01234567000100"
    assert b["spread_remessa_pct"] is None and b["spread_remessa_texto"] == "padrão"
    assert b["vigencia_fim"] == date(2026, 9, 30) and b["vigencia_inicio"] is None


def test_vigencia_desconhecida_falha_com_linha(tmp_path):
    caminho = _planilha(tmp_path, [["Loja", "12345678000100", None, None, None, None, None, None, None, None,
                                    "jan a dez", None]])
    with pytest.raises(ValueError, match="linha 4.*vigência"):
        ler_taxas(caminho)


def test_sem_aba_vigente_falha(tmp_path):
    wb = openpyxl.Workbook()
    wb.active.title = "ANTIGO 2025 - nao usar"
    caminho = tmp_path / "taxas_teste.xlsx"
    wb.save(caminho)
    with pytest.raises(ValueError, match="Taxas vigentes"):
        ler_taxas(caminho)
