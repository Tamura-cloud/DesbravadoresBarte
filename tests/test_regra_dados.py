"""Erro de dado no nível do item (R44): em cópias temporárias de dados/2026-09, um dado que impede a regra tira só o
lojista ou a linha afetada (PEN-REGRA), o resto fecha e a CLI sai com 3. Sem valores de referência: cada caso é
comparado com o fechamento completo."""

import csv
import io
import re
import shutil
from decimal import Decimal
from email import policy
from email.parser import BytesParser
from pathlib import Path

import openpyxl
import pytest

from tarvo import cli, fechamento
from tarvo.fechamento import fechar
from tarvo.leitores.orbe import ler_orbe_arquivos
from tarvo.saida import ARQUIVOS_ITENS, COMISSOES, PENDENCIAS, ItemSaida, gravar_saida, validar_saida

RAIZ = Path(__file__).resolve().parents[1]
ENTRADA = RAIZ / "dados" / "2026-09"
BRIDGE, EC_BRIDGE = "22224579000109", "1075624862"
KIWI, EC_KIWI = "07752116000126", "1038867654"
CADASTRO = "Comercial e Financeiro (Helena)"


def _ler(pasta: Path, nome: str) -> list[dict]:
    with open(pasta / nome, encoding="utf-8", newline="") as f:
        return list(csv.DictReader(f, delimiter=";"))


def _itens(pasta: Path) -> dict[str, dict]:
    return {l["id_saida"]: l for e in ARQUIVOS_ITENS for l in _ler(pasta, e.nome)}


def _regras(pasta: Path) -> dict[str, dict]:
    return {i: l for i, l in _itens(pasta).items() if i.startswith("PEN-REGRA-")}


def _um(pasta: Path, padrao: str) -> Path:
    (arquivo,) = pasta.glob(padrao)
    return arquivo


@pytest.fixture(scope="module")
def completo(tmp_path_factory):
    raiz = tmp_path_factory.mktemp("completo")
    fechar(ENTRADA, raiz / "saida", "2026-09", usar_ia=False, contestacoes=False)
    return raiz / "saida"


@pytest.fixture
def copia(tmp_path):
    pasta = tmp_path / "entrada"
    shutil.copytree(ENTRADA, pasta)
    return pasta


def _rodar(copia: Path, tmp_path: Path, capsys) -> tuple[int, str, Path]:
    saida = tmp_path / "saida"
    codigo = cli.main(["fechar", "--competencia", "2026-09", "--entrada", str(copia), "--saida", str(saida),
                       "--sem-ia", "--sem-contestacao"])
    texto = capsys.readouterr().out
    assert validar_saida(saida) == []
    return codigo, texto, saida


def _so_mudou(saida: Path, completo: Path, fora: set[str]) -> None:
    """Fora os itens afetados e as PEN-REGRA, todo item do completo continua, com os mesmos campos."""
    novos, ref = _itens(saida), _itens(completo)
    for i, l in ref.items():
        if i in fora:
            assert i not in novos, i
        else:
            assert {k: v for k, v in novos[i].items() if k != "situacao"} == l, i


# ---------------------------------------------------------------- planilha de taxas (cópia temporária)

def _planilha(copia: Path, alterar) -> None:
    caminho = _um(copia, "*taxas*.xlsx")
    wb = openpyxl.load_workbook(caminho)
    alterar(wb["Taxas vigentes"])
    wb.save(caminho)


def _linha_da_bridge_nova(ws):
    return next(r for r in ws.iter_rows(min_row=6) if r[0].value and "NOVA" in str(r[0].value))


def _taxa_orbe_do_ec(ec: str) -> tuple[int, Decimal]:
    orbe, _ = ler_orbe_arquivos(sorted(ENTRADA.glob("*orbe*.csv")))
    do_ec = orbe[orbe["cod_estabelecimento"] == ec]
    return len(do_ec), sum(do_ec["valor_taxa"], Decimal("0.00"))


def _ids_do_lojista(completo: Path, cnpj: str) -> set[str]:
    return {i for i, l in _itens(completo).items()
            if l.get("cnpj") == cnpj or l.get("tomador_cnpj") == cnpj or i.endswith(f"-{cnpj}")}


def test_duas_linhas_vigentes_para_a_bridge_na_mesma_data(completo, copia, tmp_path, capsys):
    _planilha(copia, lambda ws: setattr(_linha_da_bridge_nova(ws)[10], "value", "a partir de 10/09/2026"))
    codigo, texto, saida = _rodar(copia, tmp_path, capsys)
    assert codigo == 3
    assert "FECHAMENTO PARCIAL — fontes bloqueadas: nenhuma; itens fora por erro de dado: 1 pendência(s) " \
           f"(PEN-REGRA-VIGENCIA-CHK-{BRIDGE})" in texto
    pen = _regras(saida)[f"PEN-REGRA-VIGENCIA-CHK-{BRIDGE}"]
    assert pen["tipo"] == "ITEM_BLOQUEADO" and pen["quem_decide"] == CADASTRO
    assert "esperada 1 linha vigente na planilha de taxas em 2026-09-1" in pen["motivo"] and "achei 2 (linha 6, linha 7)" in pen["motivo"]
    n, cobrado = _taxa_orbe_do_ec(EC_BRIDGE)
    assert Decimal(pen["valor_brl"]) == cobrado and len(pen["ref_origem"].split("|")) == n
    assert f"as {n} transações do lojista (CNPJ {BRIDGE}) ficaram fora de comissão, nota e divergência" in pen["motivo"]
    assert "próximo passo: o Comercial deixa exatamente uma linha vigente por lojista e data" in pen["motivo"]
    _so_mudou(saida, completo, _ids_do_lojista(completo, BRIDGE) | {i for i in _itens(completo) if i.startswith("DIV-CHK")
                                                                   and _itens(completo)[i]["cnpj"] == BRIDGE})


def test_ec_sem_cnpj_na_planilha(completo, copia, tmp_path, capsys):
    # sem a linha da Kiwi, o EC de ec_sem_codigo (R07) aponta para um CNPJ que não está na planilha
    def tirar_kiwi(ws):
        linha = next(r[0].row for r in ws.iter_rows(min_row=6) if r[1].value and str(r[1].value) == KIWI)
        ws.delete_rows(linha)
    _planilha(copia, tirar_kiwi)
    codigo, _, saida = _rodar(copia, tmp_path, capsys)
    assert codigo == 3
    regras = _regras(saida)
    pen = regras[f"PEN-REGRA-CADASTRO-CHK-{EC_KIWI}"]
    assert f"EC {EC_KIWI}: CNPJ {KIWI} de ec_sem_codigo não está na planilha" in pen["motivo"]
    assert pen["quem_decide"] == CADASTRO and f"COM-CHK-{KIWI}" not in _itens(saida)
    # os favorecidos e pagadores com o CNPJ da Kiwi seguem na lista, sem a conferência com a planilha
    assert f"PEN-REGRA-CADASTRO-LISTA-favorecidos-{KIWI}" in regras


def test_taxa_vazia_na_planilha(completo, copia, tmp_path, capsys):
    _planilha(copia, lambda ws: setattr(_linha_da_bridge_nova(ws)[4], "value", None))
    codigo, _, saida = _rodar(copia, tmp_path, capsys)
    assert codigo == 3
    pen = _regras(saida)[f"PEN-REGRA-TAXA-CHK-{BRIDGE}"]
    assert "sem taxa_credito_vista na linha 7 da planilha de taxas" in pen["motivo"]
    assert "o Comercial informa a taxa que falta" in pen["motivo"] and pen["quem_decide"] == CADASTRO


def test_linha_da_planilha_ilegivel_tira_so_aquele_lojista(completo, copia, tmp_path, capsys):
    _planilha(copia, lambda ws: setattr(_linha_da_bridge_nova(ws)[10], "value", "a combinar"))
    codigo, _, saida = _rodar(copia, tmp_path, capsys)
    assert codigo == 3
    regras = _regras(saida)
    assert "vigência em formato desconhecido: 'a combinar'" in regras["PEN-REGRA-LINHA-TAXAS-linha_7"]["motivo"]
    assert regras["PEN-REGRA-LINHA-TAXAS-linha_7"]["quem_decide"] == CADASTRO
    assert f"a linha do CNPJ {BRIDGE} na planilha de taxas não pôde ser lida" in regras[f"PEN-REGRA-CADASTRO-CHK-{BRIDGE}"]["motivo"]
    assert f"COM-CHK-{BRIDGE}" not in _itens(saida) and f"COM-CHK-{KIWI}" in _itens(saida)


# ---------------------------------------------------------------- Orbe

def _trocar_linha_orbe(copia: Path, achar, trocar) -> tuple[str, int, str]:
    """Troca a primeira linha do relatório que casa com achar; devolve (arquivo, linha física, id_transacao)."""
    relatorio = sorted(copia.glob("*orbe*.csv"))[0]
    linhas = relatorio.read_bytes().decode("utf-8").split("\r\n")
    i = next(i for i, l in enumerate(linhas) if i > 0 and achar(l))
    linhas[i] = trocar(linhas[i])
    relatorio.write_bytes("\r\n".join(linhas).encode("utf-8"))
    return relatorio.name, i + 1, linhas[i].split(";")[0]


def test_faixa_de_parcelas_sem_taxa_tira_so_a_transacao(completo, copia, tmp_path, capsys):
    _, _, ref = _trocar_linha_orbe(copia, lambda l: ";CREDITO PARCELADO;" in l,
                                   lambda l: re.sub(r";CREDITO PARCELADO;\d+;", ";CREDITO PARCELADO;13;", l))
    codigo, _, saida = _rodar(copia, tmp_path, capsys)
    assert codigo == 3
    pen = _regras(saida)[f"PEN-REGRA-TAXA-CHK-{ref}"]
    assert f"{ref}: CREDITO PARCELADO com 13 parcelas fora das faixas de config/checkout.yaml" in pen["motivo"]
    assert "o checkout do lojista (CNPJ" in pen["motivo"] and "saiu PROVISÓRIO" in pen["motivo"]
    comissoes = {l["id_saida"]: l for l in _ler(saida, COMISSOES.nome)}
    provisorios = [i for i, l in comissoes.items() if l["situacao"] == "PROVISORIO"]
    assert len(provisorios) == 1 and provisorios[0].startswith("COM-CHK-")
    assert list(comissoes[provisorios[0]])[-1] == "situacao"  # a coluna extra é a última


def test_data_invalida_numa_linha_da_orbe(completo, copia, tmp_path, capsys):
    arquivo, n, ref = _trocar_linha_orbe(copia, lambda l: l.split(";")[1] == "02/09/2026",
                                         lambda l: l.replace(";02/09/2026;", ";31/02/2026;", 1))
    codigo, texto, saida = _rodar(copia, tmp_path, capsys)
    assert codigo == 3 and "fontes bloqueadas: nenhuma" in texto
    pen = _regras(saida)[f"PEN-REGRA-LINHA-ORBE-{arquivo}-L{n}"]
    assert f"{arquivo}, linha física {n}:" in pen["motivo"] and "31" in pen["motivo"]  # o texto do erro de data é o do Python
    assert pen["ref_origem"] == ref and pen["quem_decide"] == "Financeiro (Helena)"
    assert "PROVISÓRIO" in pen["motivo"]
    assert any(l["situacao"] == "PROVISORIO" for l in _ler(saida, COMISSOES.nome))


# ---------------------------------------------------------------- Ponte (anexo do e-mail, cópia temporária)

def _ponte(copia: Path, alterar) -> None:
    eml = _um(copia, "*.eml")
    msg = BytesParser(policy=policy.default).parse(eml.open("rb"))
    (anexo,) = list(msg.iter_attachments())
    wb = openpyxl.load_workbook(io.BytesIO(anexo.get_content()))
    alterar(wb["Operações"])
    novo = io.BytesIO()
    wb.save(novo)
    anexo.set_content(novo.getvalue(), maintype="application",
                      subtype="vnd.openxmlformats-officedocument.spreadsheetml.sheet", filename=anexo.get_filename())
    eml.write_bytes(msg.as_bytes(policy=policy.default))


def _linha_ponte(ws, ref: str) -> int:
    return next(r[0].row for r in ws.iter_rows() if r[0].value == ref)


def test_operacao_repetida_com_valores_diferentes(completo, copia, tmp_path, capsys):
    def repetir(ws):
        n = _linha_ponte(ws, "PC-0926-0003")
        valores = [c.value for c in ws[n]]
        ws.insert_rows(n + 1)
        for j, v in enumerate(valores, start=1):
            ws.cell(row=n + 1, column=j, value=v)
        ws.cell(row=n + 1, column=11, value=valores[10] + 10)  # Spread R$ diferente
    _ponte(copia, repetir)
    codigo, _, saida = _rodar(copia, tmp_path, capsys)
    assert codigo == 3
    regras = _regras(saida)
    dup = regras["PEN-REGRA-DUPLICADA-REM-PC-0926-0003"]
    assert "Nº Operação PC-0926-0003 repetido com valores diferentes nas linhas 7 e 8 (colunas: spread_brl)" in dup["motivo"]
    assert dup["valor_brl"] == "0.00" and "o fechamento não escolhe" in dup["motivo"]
    assert "PEN-REGRA-LINHA-CONF-PC-0926-0003" in regras  # fora também das conferências P11 e P12
    nota = next(l for l in _ler(saida, "notas_a_emitir.csv") if l["id_saida"].startswith("NF-REM-"))
    assert nota["situacao"] == "PROVISORIO"


def test_operacao_sem_data_de_contratacao(completo, copia, tmp_path, capsys):
    _ponte(copia, lambda ws: setattr(ws.cell(row=_linha_ponte(ws, "PC-0926-0003"), column=2), "value", None))
    codigo, _, saida = _rodar(copia, tmp_path, capsys)
    assert codigo == 3
    regras = _regras(saida)
    assert "PC-0926-0003 (linha 7 do relatório da Ponte): data_contratacao vazio" in regras["PEN-REGRA-LINHA-REM-PC-0926-0003"]["motivo"]
    assert "PEN-REGRA-LINHA-CONF-PC-0926-0003" in regras


def test_cnpj_invalido_numa_linha_da_ponte(completo, copia, tmp_path, capsys):
    _ponte(copia, lambda ws: setattr(ws.cell(row=_linha_ponte(ws, "PC-0926-0003"), column=4), "value", "123"))
    codigo, _, saida = _rodar(copia, tmp_path, capsys)
    assert codigo == 3
    pen = _regras(saida)["PEN-REGRA-LINHA-PONTE_EML-linha_7"]
    assert "linha 7: CNPJ inválido: '123'" in pen["motivo"] and pen["ref_origem"] == "PC-0926-0003"


# ---------------------------------------------------------------- extratos

def test_valor_invalido_numa_linha_do_aurora(completo, copia, tmp_path, capsys):
    aurora = _um(copia, "*aurora*.csv")
    aurora.write_bytes(aurora.read_bytes().replace(b";-389,00;", b";-38x,00;", 1))
    codigo, _, saida = _rodar(copia, tmp_path, capsys)
    assert codigo == 3
    pen = _regras(saida)["PEN-REGRA-LINHA-AURORA-L10"]
    assert "L10: número em formato brasileiro inválido: '-38x,00'" in pen["motivo"]
    _so_mudou(saida, completo, {"LAN-AUR-L10"})


def test_fitid_repetido_no_meridiano_tira_as_duas_copias(completo, copia, tmp_path, capsys):
    ofx = _um(copia, "*.ofx")
    texto = ofx.read_bytes().decode("cp1252")
    bloco = re.search(r"<STMTTRN>.*?</STMTTRN>\s*", texto, re.S).group(0)
    fitid = re.search(r"<FITID>([^<\r\n]+)", bloco).group(1).strip()
    ofx.write_bytes(texto.replace(bloco, bloco + bloco, 1).encode("cp1252"))
    codigo, _, saida = _rodar(copia, tmp_path, capsys)
    assert codigo == 3
    regras = _regras(saida)
    assert {f"PEN-REGRA-LINHA-MERIDIANO-{fitid}", f"PEN-REGRA-LINHA-MERIDIANO-{fitid}-2"} <= set(regras)
    assert f"LAN-MER-{fitid}" not in _itens(saida)


def test_meridiano_sem_periodo(completo, copia, tmp_path, capsys):
    ofx = _um(copia, "*.ofx")
    ofx.write_bytes(re.sub(rb"<DTSTART>[^<\r\n]*", b"", ofx.read_bytes(), count=1))
    codigo, _, saida = _rodar(copia, tmp_path, capsys)
    assert codigo == 3
    pen = _regras(saida)["PEN-REGRA-PERIODO-MERIDIANO"]
    assert "sem DTSTART/DTEND" in pen["motivo"] and "P11" in pen["motivo"]
    assert not [i for i in _itens(saida) if i.startswith("PEN-TRANSF-")]  # P13 não foi feita
    assert [i for i in _itens(saida) if i.startswith("LAN-MER-")]


# ---------------------------------------------------------------- erro de programa continua parando

def test_erro_de_programa_no_leitor_nao_vira_bloqueio(copia, tmp_path, monkeypatch):
    def quebra(*_):
        raise KeyError("bug no leitor")
    monkeypatch.setitem(fechamento.LEITORES, "aurora", quebra)
    with pytest.raises(KeyError):
        cli.main(["fechar", "--competencia", "2026-09", "--entrada", str(copia), "--saida", str(tmp_path / "s")])
    assert not (tmp_path / "s").exists()


def test_situacao_e_a_ultima_coluna_mesmo_com_outra_extra(tmp_path):
    item = ItemSaida("PEN-X", PENDENCIAS.nome, {"situacao": "PROVISORIO", "tipo": "T", "ref_origem": "R",
                                                 "motivo": "m", "quem_decide": "q", "outra": "x"},
                     Decimal("0.00"), [("a.csv", "R", Decimal("0.00"))])
    gravar_saida(tmp_path, [item])
    assert (tmp_path / PENDENCIAS.nome).read_text(encoding="utf-8").splitlines()[0].endswith(";outra;situacao")
    assert validar_saida(tmp_path) == []
